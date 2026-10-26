# -*- coding: utf-8 -*-
"""Krone Mittagstisch -> Word
Inserisci il link della pagina: il programma legge il Mittagstisch del giorno
e salva un file Word sempre con lo stesso layout (font, intestazione, colori).
"""
import os
import re
import sys
import time
import threading
import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: F401
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm, Twips

# ---------- Impostazioni di stile (uguali ogni giorno) ----------
FONT = "Quicksand"                 # font del sito
GREEN = RGBColor(0x07, 0x33, 0x34)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
TEXT = RGBColor(0x1A, 0x1A, 0x1A)
DEFAULT_URL = "https://trattoria-pizzeria-krone.de/home#a9be7e1c-5c87-4fd2-88f6-b17fd602523c"

DAYS = ["Montag", "Dienstag", "Mittwoch", "Donnerstag", "Freitag", "Samstag", "Sonntag"]
PRICE_RE = re.compile(r"^\s*(?:€\s*)?(\d{1,3}(?:[.,]\d{1,2})?)\s*(?:€|EUR)?\s*$")


# ---------- Lettura del testo ----------
def parse_day(text, day):
    """Restituisce [(piatto, prezzo), ...] per il giorno indicato."""
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    days_low = {d.lower() for d in DAYS}
    for i, line in enumerate(lines):
        if line.lower() != day.lower():
            continue
        dishes, buf = [], []
        for x in lines[i + 1:]:
            if x.lower() in days_low:
                break
            m = PRICE_RE.match(x)
            if m:
                if buf:
                    dishes.append((" ".join(buf[-2:]), m.group(1)))
                buf = []
            else:
                buf.append(x)
                if dishes and len(buf) > 3:   # siamo fuori dalla sezione
                    break
        if dishes:
            return dishes
    return []


def fetch_text_and_parse(url, day, log=lambda s: None):
    """Apre la pagina con Edge (già presente su Windows) e legge il testo."""
    from selenium import webdriver
    from selenium.webdriver.edge.options import Options

    opts = Options()
    opts.add_argument("--headless=new")
    opts.add_argument("--window-size=1280,2400")
    opts.add_argument("--lang=de-DE")
    driver = webdriver.Edge(options=opts)
    try:
        driver.get(url)
        text, dishes = "", []
        for _ in range(25):               # fino a ~25 secondi
            time.sleep(1)
            text = driver.execute_script("return document.body.innerText") or ""
            dishes = parse_day(text, day)
            if dishes:
                break
        return dishes, text
    finally:
        driver.quit()


# ---------- Creazione del Word ----------
def _set_font(run, size, color=TEXT, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), FONT)


def _bottom_border(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "12")
    b.set(qn("w:space"), "8")
    b.set(qn("w:color"), "C9A227")
    borders.append(b)
    ppr.append(borders)


def build_docx(day, dishes, path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    width_tw = int((21.0 - 4.6) / 2.54 * 1440)

    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(13)

    # Intestazione (sempre uguale)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("KRONE")
    _set_font(r, 14, GOLD, True)
    rpr = r._element.get_or_add_rPr()
    sp = OxmlElement("w:spacing"); sp.set(qn("w:val"), "120"); rpr.append(sp)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    _set_font(p.add_run("Mittagstisch"), 32, GREEN)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    _set_font(p.add_run("Trattoria · Pizzeria"), 11, GREEN)
    _bottom_border(p)

    # Contenuto (cambia ogni giorno)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    _set_font(p.add_run(day), 22, GREEN)

    for name, price in dishes:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.keep_together = True
        p.paragraph_format.tab_stops.add_tab_stop(
            Twips(width_tw), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        _set_font(p.add_run(name), 13, TEXT)
        _set_font(p.add_run("\t" + price + " €"), 13, GREEN, True)

    doc.save(path)


# ---------- Cartella di salvataggio ----------
def output_dir():
    base = os.path.dirname(sys.executable if getattr(sys, "frozen", False) else os.path.abspath(__file__))
    d = os.path.join(base, "Mittagstisch")
    os.makedirs(d, exist_ok=True)
    return d


def make_file(url, day, log=lambda s: None):
    dishes, text = fetch_text_and_parse(url, day, log)
    if not dishes:
        dbg = os.path.join(output_dir(), "debug_testo_pagina.txt")
        with open(dbg, "w", encoding="utf-8") as f:
            f.write(text)
        raise RuntimeError(
            f"Nessun piatto trovato per {day}.\nHo salvato il testo letto dalla pagina in:\n{dbg}\n"
            "Mandalo a Claude per sistemare la lettura.")
    stamp = datetime.date.today().strftime("%Y-%m-%d")
    path = os.path.join(output_dir(), f"Mittagstisch_{day}_{stamp}.docx")
    build_docx(day, dishes, path)
    return path, len(dishes)


# ---------- Finestra ----------
def main():
    import tkinter as tk
    from tkinter import ttk, messagebox

    root = tk.Tk()
    root.title("Krone – Mittagstisch in Word")
    root.geometry("560x230")
    root.configure(bg="#FFFFFF")
    pad = {"padx": 16, "pady": 6}

    tk.Label(root, text="Link della pagina:", bg="#FFFFFF", fg="#073334",
             font=("Segoe UI", 10, "bold")).pack(anchor="w", **pad)
    url_var = tk.StringVar(value=DEFAULT_URL)
    ttk.Entry(root, textvariable=url_var).pack(fill="x", padx=16)

    row = tk.Frame(root, bg="#FFFFFF"); row.pack(fill="x", **pad)
    tk.Label(row, text="Giorno:", bg="#FFFFFF", fg="#073334",
             font=("Segoe UI", 10, "bold")).pack(side="left")
    today = DAYS[datetime.date.today().weekday()]
    day_var = tk.StringVar(value=today if today in DAYS[:5] else "Montag")
    ttk.Combobox(row, textvariable=day_var, values=DAYS, state="readonly",
                 width=14).pack(side="left", padx=8)

    status = tk.StringVar(value="Pronto.")
    tk.Label(root, textvariable=status, bg="#FFFFFF", fg="#444444",
             wraplength=520, justify="left").pack(anchor="w", padx=16, pady=4)

    btn = tk.Button(root, text="Crea Word", bg="#073334", fg="white",
                    font=("Segoe UI", 11, "bold"), relief="flat", padx=18, pady=6)
    btn.pack(pady=6)

    def work():
        try:
            path, n = make_file(url_var.get().strip(), day_var.get())
            root.after(0, lambda: done(path, n, None))
        except Exception as e:  # noqa: BLE001
            root.after(0, lambda: done(None, 0, e))

    def done(path, n, err):
        btn.config(state="normal")
        if err:
            status.set("Errore.")
            messagebox.showerror("Errore", str(err))
        else:
            status.set(f"Fatto: {n} piatti. File: {os.path.basename(path)}")
            try:
                os.startfile(path)  # apre il Word
            except Exception:  # noqa: BLE001
                pass

    def start():
        if not url_var.get().strip():
            messagebox.showwarning("Link mancante", "Incolla il link della pagina.")
            return
        btn.config(state="disabled")
        status.set("Leggo la pagina… (qualche secondo)")
        threading.Thread(target=work, daemon=True).start()

    btn.config(command=start)
    root.mainloop()


if __name__ == "__main__":
    main()
